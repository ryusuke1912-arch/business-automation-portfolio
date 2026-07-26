# -*- coding: utf-8 -*-
"""
Webスクレイピング → CSV / Excel 納品ツール（ポートフォリオ公開版）

対象 : https://books.toscrape.com/
       スクレイピング学習のために公開されている練習用サイト。
       実サービスへ無断でクローラを当てないため、デモの対象はここに固定している。

実行 : python scraper.py --pages 2
出力 : out/books.csv （Excelで開いても文字化けしないUTF-8 BOM付き）
       out/books.xlsx（一覧＋カテゴリ別集計）
       out/scrape.log

■ 「取れればいい」で終わらせないための実装
  1. robots.txt を取得して対象URLが許可されているか確認し、禁止なら中止する
  2. User-Agent に用途と連絡先を明示する（無記名のbotとして扱われないため）
  3. リクエスト間隔を空け、相手サーバに負荷をかけない
  4. 失敗時は指数バックオフで再試行。3回失敗した行は欠損として記録し全体は止めない
  5. 取得済みHTMLをキャッシュし、中断しても続きから再開できる
  6. 件数・欠損の検算結果をログに残し、納品物が信用できる状態にする
"""
import os, csv, re, time, json, argparse, hashlib, random
import urllib.parse as up
import urllib.robotparser as urobot
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

BASE = 'https://books.toscrape.com/'
START = up.urljoin(BASE, 'catalogue/page-1.html')
UA = ('PortfolioScraper/1.0 (educational demo; contact: example@example.com) '
      'python-requests')

HERE = os.path.dirname(os.path.abspath(__file__))
OUTDIR = os.path.join(HERE, 'out')
CACHE = os.path.join(HERE, '.cache')
os.makedirs(OUTDIR, exist_ok=True)
os.makedirs(CACHE, exist_ok=True)
LOG = os.path.join(OUTDIR, 'scrape.log')

RATING = {'One': 1, 'Two': 2, 'Three': 3, 'Four': 4, 'Five': 5}


def log(msg):
    line = f'[{time.strftime("%H:%M:%S")}] {msg}'
    print(line)
    with open(LOG, 'a', encoding='utf-8') as f:
        f.write(line + '\n')


# ============================================================ 取得層
class Fetcher:
    def __init__(self, delay=0.4, retries=3, use_cache=True):
        self.s = requests.Session()
        self.s.headers.update({'User-Agent': UA,
                               'Accept-Language': 'en-US,en;q=0.9'})
        self.delay = delay
        self.retries = retries
        self.use_cache = use_cache
        self.n_net = 0
        self.n_cache = 0
        self.rp = urobot.RobotFileParser()
        try:
            self.rp.set_url(up.urljoin(BASE, '/robots.txt'))
            self.rp.read()
            self.robots_ok = True
        except Exception as e:
            log(f'robots.txt を取得できませんでした（{e}）。安全側に倒して中止します。')
            self.robots_ok = False

    def allowed(self, url):
        if not self.robots_ok:
            return False
        try:
            return self.rp.can_fetch(UA, url)
        except Exception:
            return False

    def _cpath(self, url):
        return os.path.join(CACHE, hashlib.sha1(url.encode()).hexdigest() + '.html')

    def get(self, url):
        """robots確認 → キャッシュ → ネットワーク（指数バックオフ付き）"""
        if not self.allowed(url):
            log(f'robots.txt により許可されていないためスキップ: {url}')
            return None
        cp = self._cpath(url)
        if self.use_cache and os.path.exists(cp):
            self.n_cache += 1
            with open(cp, encoding='utf-8') as f:
                return f.read()
        wait = 1.0
        for attempt in range(1, self.retries + 1):
            try:
                time.sleep(self.delay + random.uniform(0, 0.2))   # 一定間隔＋ゆらぎ
                r = self.s.get(url, timeout=20)
                self.n_net += 1
                if r.status_code == 200:
                    r.encoding = r.apparent_encoding or 'utf-8'
                    with open(cp, 'w', encoding='utf-8') as f:
                        f.write(r.text)
                    return r.text
                if r.status_code == 404:
                    log(f'404 のため打ち切り: {url}')
                    return None
                log(f'HTTP {r.status_code} ({attempt}/{self.retries}): {url}')
            except requests.RequestException as e:
                log(f'通信エラー ({attempt}/{self.retries}): {type(e).__name__}: {url}')
            time.sleep(wait)
            wait *= 2                                             # 指数バックオフ
        log(f'★ {self.retries}回失敗したため欠損として記録: {url}')
        return None


# ============================================================ 解析層
def money(txt):
    m = re.search(r'([\d.]+)', (txt or '').replace(',', ''))
    return float(m.group(1)) if m else None


def parse_list(html, page_url):
    """一覧ページ → 詳細URLのリスト と 次ページURL"""
    soup = BeautifulSoup(html, 'html.parser')
    links = []
    for a in soup.select('article.product_pod h3 a'):
        links.append(up.urljoin(page_url, a.get('href')))
    nxt = soup.select_one('li.next a')
    return links, (up.urljoin(page_url, nxt.get('href')) if nxt else None)


def parse_detail(html, url):
    soup = BeautifulSoup(html, 'html.parser')
    rec = {'URL': url}

    h1 = soup.select_one('div.product_main h1')
    rec['書名'] = h1.get_text(strip=True) if h1 else None

    p = soup.select_one('div.product_main p.price_color')
    rec['価格(GBP)'] = money(p.get_text()) if p else None

    star = soup.select_one('div.product_main p.star-rating')
    rec['評価'] = None
    if star:
        for c in star.get('class', []):
            if c in RATING:
                rec['評価'] = RATING[c]

    av = soup.select_one('div.product_main p.availability')
    txt = av.get_text(' ', strip=True) if av else ''
    m = re.search(r'(\d+)\s*available', txt)
    rec['在庫数'] = int(m.group(1)) if m else None
    rec['在庫状態'] = 'In stock' if 'In stock' in txt else ('Out of stock' if txt else None)

    crumbs = [a.get_text(strip=True) for a in soup.select('ul.breadcrumb li a')]
    rec['カテゴリ'] = crumbs[-1] if len(crumbs) >= 3 else None

    # 商品情報テーブル（UPC・税・レビュー数）
    info = {}
    for tr in soup.select('table.table-striped tr'):
        th, td = tr.find('th'), tr.find('td')
        if th and td:
            info[th.get_text(strip=True)] = td.get_text(strip=True)
    rec['UPC'] = info.get('UPC')
    rec['税(GBP)'] = money(info.get('Tax'))
    rec['レビュー数'] = int(info['Number of reviews']) if info.get('Number of reviews', '').isdigit() else None

    d = soup.select_one('#product_description ~ p')
    desc = d.get_text(' ', strip=True) if d else ''
    rec['説明(抜粋)'] = (desc[:120] + '…') if len(desc) > 120 else desc
    return rec


# ============================================================ 出力層
COLS = ['書名', 'カテゴリ', '価格(GBP)', '税(GBP)', '評価', '在庫状態', '在庫数',
        'レビュー数', 'UPC', '説明(抜粋)', 'URL']
BOLD = Font(bold=True)
HEAD = PatternFill('solid', fgColor='FFD9E1F2')
MISS = PatternFill('solid', fgColor='FFFFF2CC')
thin = Side(style='thin', color='FFBFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_csv(rows):
    path = os.path.join(OUTDIR, 'books.csv')
    # utf-8-sig … Excelでダブルクリックしても文字化けしない
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=COLS, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows)
    return path


def write_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '一覧'
    for j, c in enumerate(COLS, 1):
        cc = ws.cell(row=1, column=j, value=c)
        cc.font, cc.fill, cc.border = BOLD, HEAD, BD
        cc.alignment = Alignment(horizontal='center')
    for i, r in enumerate(rows, 2):
        for j, c in enumerate(COLS, 1):
            cc = ws.cell(row=i, column=j, value=r.get(c))
            cc.border = BD
            if r.get(c) is None:
                cc.fill = MISS                        # 欠損は塗って可視化
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{chr(64+len(COLS))}{len(rows)+1}'
    for col, w in zip('ABCDEFGHIJK', (44, 18, 11, 10, 7, 12, 9, 11, 18, 46, 60)):
        ws.column_dimensions[col].width = w

    agg = {}
    for r in rows:
        k = r.get('カテゴリ') or '(不明)'
        a = agg.setdefault(k, {'カテゴリ': k, '冊数': 0, '価格計': 0.0, '評価計': 0, '評価数': 0})
        a['冊数'] += 1
        a['価格計'] += r.get('価格(GBP)') or 0
        if r.get('評価'):
            a['評価計'] += r['評価']
            a['評価数'] += 1
    ws2 = wb.create_sheet('カテゴリ別')
    heads = ['カテゴリ', '冊数', '平均価格(GBP)', '平均評価']
    for j, c in enumerate(heads, 1):
        cc = ws2.cell(row=1, column=j, value=c)
        cc.font, cc.fill, cc.border = BOLD, HEAD, BD
    for i, a in enumerate(sorted(agg.values(), key=lambda x: -x['冊数']), 2):
        ws2.cell(row=i, column=1, value=a['カテゴリ'])
        ws2.cell(row=i, column=2, value=a['冊数'])
        ws2.cell(row=i, column=3, value=round(a['価格計'] / a['冊数'], 2))
        ws2.cell(row=i, column=4,
                 value=round(a['評価計'] / a['評価数'], 2) if a['評価数'] else None)
    for col, w in zip('ABCD', (22, 8, 15, 10)):
        ws2.column_dimensions[col].width = w

    path = os.path.join(OUTDIR, 'books.xlsx')
    wb.save(path)
    return path


# ============================================================ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--pages', type=int, default=2, help='巡回する一覧ページ数')
    ap.add_argument('--delay', type=float, default=0.4, help='リクエスト間隔(秒)')
    ap.add_argument('--no-cache', action='store_true', help='キャッシュを使わず取り直す')
    a = ap.parse_args()

    open(LOG, 'w', encoding='utf-8').close()
    t0 = time.time()
    fx = Fetcher(delay=a.delay, use_cache=not a.no_cache)
    if not fx.robots_ok:
        return
    log(f'robots.txt 確認OK / 開始 pages={a.pages} delay={a.delay}s')

    # ---- 一覧を辿って詳細URLを集める ----
    urls, page, n = [], START, 0
    while page and n < a.pages:
        html = fx.get(page)
        if not html:
            break
        links, page = parse_list(html, page)
        urls += links
        n += 1
        log(f'一覧 {n}ページ目: {len(links)}件 （累計 {len(urls)}件）')

    # ---- 詳細を取得 ----
    rows, failed = [], 0
    for i, u in enumerate(urls, 1):
        html = fx.get(u)
        if not html:
            failed += 1
            rows.append({'URL': u})
            continue
        rows.append(parse_detail(html, u))
        if i % 10 == 0 or i == len(urls):
            log(f'詳細 {i}/{len(urls)} 件取得')

    cpath, xpath = write_csv(rows), write_xlsx(rows)

    # ---- 検算 ----
    filled = sum(1 for r in rows for c in COLS if r.get(c) is not None)
    cells = len(rows) * len(COLS)
    missing_cols = {c: sum(1 for r in rows if r.get(c) is None) for c in COLS}
    dup = len(rows) - len({r.get('UPC') for r in rows if r.get('UPC')}) - failed

    log('=' * 52)
    log(f'  取得件数        : {len(rows)} 件（取得失敗 {failed} 件）')
    log(f'  項目の充足率    : {filled/cells*100:.1f}%  ({filled}/{cells} セル)')
    log(f'  UPC重複         : {max(dup,0)} 件')
    log(f'  ネットワーク要求: {fx.n_net} 回 / キャッシュ利用 {fx.n_cache} 回')
    log(f'  所要時間        : {time.time()-t0:.1f} 秒')
    log('=' * 52)
    for c, k in missing_cols.items():
        if k:
            log(f'  欠損あり: {c} … {k}件')
    log(f'出力: {cpath}')
    log(f'出力: {xpath}')


if __name__ == '__main__':
    main()
