# -*- coding: utf-8 -*-
"""
個別指導・シフト 自動配置エンジン（ポートフォリオ公開版）

入力 : students.csv / teachers.csv / dates.csv
出力 : schedule_out.xlsx（日付別シート＋稼働率サマリ）
実行 : python placer.py          ※Excel本体は不要（openpyxlのみ）

■ 解いている問題
  「講師の出勤可否」「生徒ごとの受講不可枠」「1日あたりの上限コマ数」を
  すべて満たしながら、必要コマ数を期間内に詰め込む（制約充足問題）。
  さらに現場の要求として次の2つを同時に最大化する：
    (1) 1対2率  … 1講師2生徒の座席を埋めて人件費効率を上げる
    (2) 連続性  … 生徒の1日の授業に空きコマを作らない（拘束時間を減らす）
  この2つはトレードオフの関係にあるため、配置後に局所改善パスをかけて両立させる。

■ アルゴリズム
  1) 前処理 : 各「生徒×科目」について物理的に置ける枠数を数え、
              余裕(slack) = 置ける枠数 - 必要コマ数 が小さい順に処理（MRV法）。
              制約の厳しい生徒を先に通すことで、後半の詰まりを防ぐ。
  2) 配置   : 各コマを、複合スコアが最大の枠へ貪欲配置。
              スコア = (連続かつ1:2, 連続, 1:2, 隣接, 同日寄せ, 負荷分散, 日付)
  3) 改善A  : 1:1の枠どうしを同一講師の別枠へ寄せて 1:2 に統合（pair_pass）
  4) 改善B  : 生徒の空きコマを同日内移動で解消（gap_pass / 1:2を減らさない範囲で）
  5) 3→4→3 を再実行して収束させる
"""
import csv, re, os, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'schedule_out.xlsx')

SLOTNAME = ['13:00', '14:30', '17:00', '18:30', '20:00']
TIMEMAP = {t: i for i, t in enumerate(SLOTNAME)}
SEATS = 2                      # 1講師1コマあたりの定員（1対2）
WD = '月火水木金土日'


def norm(s):
    return unicodedata.normalize('NFKC', str(s)).replace(' ', '').replace('　', '') if s else s


def contiguous(s):
    return (not s) or (max(s) - min(s) + 1 == len(s))


# ============================================================ 入力
def load_all():
    dates = []
    with open(os.path.join(HERE, 'dates.csv'), encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            m, d = r['日付'].split('/')
            dates.append({'ds': r['日付'], 'md': (int(m), int(d)),
                          'wd': r['曜日'], 'sheet': r['シート名']})
    dates.sort(key=lambda x: x['md'])

    av, teachers = {}, []
    with open(os.path.join(HERE, 'teachers.csv'), encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            t = r['先生']
            if t not in teachers:
                teachers.append(t)
            av[(norm(t), r['日付'])] = {b for b, s in enumerate(SLOTNAME) if r.get(s, '').strip()}

    students, order = {}, []
    with open(os.path.join(HERE, 'students.csv'), encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            nm = r['生徒名']
            if nm not in students:
                students[nm] = {'grade': r['学年'], 'cap': int(r['1日上限'] or 2), 'ng': '', 'subs': []}
                order.append(nm)
            if r.get('NG', '').strip():
                students[nm]['ng'] = r['NG'].strip()
            tl = [x.strip() for x in r['担当(優先順;区切り)'].split(';') if x.strip()]
            students[nm]['subs'].append((r['科目'], int(r['回数']), tl,
                                         r['配分(順/分散)'].strip() or '順'))
    return dates, teachers, av, students, order


# ============================================================ NG構文
def parse_ng(text):
    """'土;月18:30以降;火17:00以降;8/1' → (全日共通, 曜日別, 日付別) のブロック集合"""
    g, wd, dt = set(), {}, {}
    for tok in re.split(r'[;；]', text or ''):
        tok = tok.strip()
        if not tok:
            continue
        w = None
        if tok[0] in WD:
            w, tok = tok[0], tok[1:]
        date = None
        m = re.match(r'(\d+)/(\d+)', tok)
        if m:
            date = (int(m.group(1)), int(m.group(2)))
            tok = tok[m.end():].strip()
        if tok == '':
            blocks = set(range(5))                                   # 終日
        else:
            mm = re.match(r'(\d{1,2}:\d{2})(以降)?', tok)
            if not mm or mm.group(1) not in TIMEMAP:
                print(f'  [NG解釈不可] {tok}')
                continue
            base = TIMEMAP[mm.group(1)]
            blocks = set(range(base, 5)) if mm.group(2) else {base}
        if date is not None:
            dt[date] = dt.get(date, set()) | blocks
        elif w is not None:
            wd[w] = wd.get(w, set()) | blocks
        else:
            g |= blocks
    return g, wd, dt


def ng_blocks(rule, md, w):
    g, wd, dt = rule
    return g | wd.get(w, set()) | dt.get(md, set())


# ============================================================ 状態
class Board:
    def __init__(self, dates, teachers, av, students):
        self.dates = dates
        self.teachers = teachers
        self.av = av
        self.students = students
        self.byd = {d['sheet']: d for d in dates}
        # seat[(sheet, teacher, block)] = [ {name,grade,subj}, ... ] 最大SEATS
        self.seat = {}
        self.busy = {n: {d['sheet']: set() for d in dates} for n in students}
        self.load = {n: {d['sheet']: 0 for d in dates} for n in students}

    def avail(self, teacher, ds, b):
        return b in self.av.get((norm(teacher), ds), set())

    def occ(self, sheet, teacher, b):
        return self.seat.get((sheet, teacher, b), [])

    def put(self, sheet, teacher, b, ent):
        self.seat.setdefault((sheet, teacher, b), []).append(ent)
        self.busy[ent['name']][sheet].add(b)
        self.load[ent['name']][sheet] += 1

    def pop(self, sheet, teacher, b, name):
        lst = self.seat.get((sheet, teacher, b), [])
        for i, e in enumerate(lst):
            if e['name'] == name:
                lst.pop(i)
                self.busy[name][sheet].discard(b)
                self.load[name][sheet] -= 1
                return e
        return None


# ============================================================ 配置
def place_one(bd, task, teacher):
    """task の1コマを teacher の最良枠へ貪欲配置。置けたら True。"""
    name, cap, rule = task['name'], task['cap'], task['rule']
    best = None
    for d in bd.dates:
        sh, md_, w = d['sheet'], d['md'], d['wd']
        if bd.load[name][sh] >= cap:
            continue
        ngb = ng_blocks(rule, md_, w)
        S = bd.busy[name][sh]
        for b in range(5):
            if b in S or b in ngb:
                continue
            if not bd.avail(teacher, d['ds'], b):
                continue
            cur = bd.occ(sh, teacher, b)
            if len(cur) >= SEATS:
                continue
            if any(e['name'] == name for e in cur):
                continue
            newset = S | {b}
            contig = 1 if contiguous(newset) else 0
            extend = 1 if (S and (b == min(S) - 1 or b == max(S) + 1)) else 0
            same_day = 1 if S else 0
            pair = 1 if len(cur) == 1 else 0            # 入れると1対2が成立
            sc = (contig + pair, contig, pair, extend, same_day,
                  -bd.load[name][sh], -(md_[0] * 100 + md_[1]))
            if best is None or sc > best[0]:
                best = (sc, sh, b)
    if best is None:
        return False
    _, sh, b = best
    bd.put(sh, teacher, b, {'name': name, 'grade': task['grade'], 'subj': task['subj']})
    return True


def feasible_slots(bd, task):
    """出勤×NGを満たす物理的な枠数（座席の埋まり具合は見ない上界）"""
    cnt = 0
    for d in bd.dates:
        ngb = ng_blocks(task['rule'], d['md'], d['wd'])
        for b in range(5):
            if b in ngb:
                continue
            if any(bd.avail(t, d['ds'], b) for t in task['teachers']):
                cnt += 1
    return cnt


# ============================================================ 改善パス
def pair_pass(bd):
    """1対1の枠を同一講師の別の1対1枠へ寄せて1対2に統合する。"""
    moved = 0
    progress = True
    while progress:
        progress = False
        singles = [(k, v[0]) for k, v in bd.seat.items() if len(v) == 1]
        idx = {}
        for (sh, t, b), e in singles:
            idx.setdefault(t, []).append((sh, b, e))
        for (sh, t, b), e in singles:
            nm = e['name']
            if nm not in bd.students:
                continue
            for (sh2, b2, e2) in idx.get(t, []):
                if (sh2, b2) == (sh, b) or e2['name'] == nm:
                    continue
                if len(bd.occ(sh2, t, b2)) != 1:
                    continue
                if not can_move(bd, nm, sh, b, sh2, b2, t):
                    continue
                bd.pop(sh, t, b, nm)
                bd.put(sh2, t, b2, e)
                moved += 1
                progress = True
                break
            if progress:
                break
    return moved


def can_move(bd, name, sh_from, b_from, sh_to, b_to, teacher):
    """name の (sh_from,b_from) のコマを (sh_to,b_to) へ動かせるか。"""
    st = bd.students[name]
    rule, cap = st['_rule'], st['cap']
    d = bd.byd[sh_to]
    if not bd.avail(teacher, d['ds'], b_to):
        return False
    if len(bd.occ(sh_to, teacher, b_to)) >= SEATS:
        return False
    if b_to in ng_blocks(rule, d['md'], d['wd']):
        return False
    # 移動元の日: 抜いた後も連続を維持
    rest_from = bd.busy[name][sh_from] - {b_from}
    if not contiguous(rest_from):
        return False
    # 移動先の日
    dest = (bd.busy[name][sh_to] - ({b_from} if sh_to == sh_from else set())) | {b_to}
    if b_to in bd.busy[name][sh_to] and not (sh_to == sh_from and b_to == b_from):
        return False
    if not contiguous(dest):
        return False
    load_to = bd.load[name][sh_to] - (1 if sh_to == sh_from else 0)
    if load_to >= cap:
        return False
    return True


def gap_pass(bd):
    """生徒の空きコマ（1日の中の非連続）を同日内移動で解消。1対2を減らす移動はしない。"""
    fixed = 0
    progress = True
    while progress:
        progress = False
        for nm in bd.students:
            for d in bd.dates:
                sh = d['sheet']
                bl = bd.busy[nm][sh]
                if len(bl) < 2 or contiguous(bl):
                    continue
                holes = sorted(set(range(min(bl), max(bl) + 1)) - bl)
                done = False
                for hole in holes:
                    for src_b in (hole - 1, hole + 1):
                        if src_b not in bl:
                            continue
                        src_t = next((t for (s2, t, b2) in bd.seat
                                      if s2 == sh and b2 == src_b
                                      and any(e['name'] == nm for e in bd.seat[(s2, t, b2)])), None)
                        if src_t is None:
                            continue
                        src_pair = 1 if len(bd.occ(sh, src_t, src_b)) == SEATS else 0
                        dst_pair = 1 if len(bd.occ(sh, src_t, hole)) == 1 else 0
                        if dst_pair - src_pair < 0:
                            continue
                        if not can_move(bd, nm, sh, src_b, sh, hole, src_t):
                            continue
                        e = bd.pop(sh, src_t, src_b, nm)
                        bd.put(sh, src_t, hole, e)
                        fixed += 1
                        progress = done = True
                        break
                    if done:
                        break
                if done:
                    break
            if progress:
                break
    return fixed


# ============================================================ 出力
FONT = Font(name='Calibri', size=10)
BOLD = Font(name='Calibri', size=10, bold=True)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin = Side(style='thin', color='FFBFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
ORANGE = PatternFill('solid', fgColor='FFFFC000')
GRAY = PatternFill('solid', fgColor='FFF2F2F2')
HEAD = PatternFill('solid', fgColor='FFD9E1F2')


def write_excel(bd, kpi, report):
    wb = Workbook()
    wb.remove(wb.active)

    for d in bd.dates:
        sh = d['sheet']
        ws = wb.create_sheet(sh)
        ws.cell(row=1, column=1, value='講師').fill = HEAD
        ws.cell(row=1, column=1).font = BOLD
        ws.cell(row=1, column=1).alignment = CEN
        for b, t in enumerate(SLOTNAME):
            c0 = 2 + b * 3
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 2)
            cc = ws.cell(row=1, column=c0, value=t)
            cc.font, cc.alignment, cc.fill, cc.border = BOLD, CEN, HEAD, BORDER
        r = 2
        for t in bd.teachers:
            ws.merge_cells(start_row=r, start_column=1, end_row=r + SEATS - 1, end_column=1)
            cc = ws.cell(row=r, column=1, value=t)
            cc.font, cc.alignment, cc.border = FONT, CEN, BORDER
            for b in range(5):
                c0 = 2 + b * 3
                occ = bd.occ(sh, t, b)
                open_seat = bd.avail(t, d['ds'], b)
                for k in range(SEATS):
                    rr = r + k
                    e = occ[k] if k < len(occ) else None
                    vals = [e['grade'], e['name'], e['subj']] if e else ['', '', '']
                    for j, v in enumerate(vals):
                        c = ws.cell(row=rr, column=c0 + j, value=v)
                        c.font, c.alignment, c.border = FONT, CEN, BORDER
                        if e:
                            c.fill = ORANGE
                        elif not open_seat:
                            c.fill = GRAY                    # 出勤なし＝灰色
            r += SEATS
        ws.freeze_panes = 'B2'
        ws.column_dimensions['A'].width = 12
        for b in range(5):
            c0 = 2 + b * 3
            ws.column_dimensions[get_column_letter(c0)].width = 5
            ws.column_dimensions[get_column_letter(c0 + 1)].width = 11
            ws.column_dimensions[get_column_letter(c0 + 2)].width = 6

    # ---- サマリシート ----
    ws = wb.create_sheet('サマリ', 0)
    ws['A1'] = '自動配置レポート'
    ws['A1'].font = Font(name='Calibri', size=13, bold=True)
    ws['A2'] = '※ 本ファイルの氏名・データはすべて架空のサンプルです'
    ws['A2'].font = Font(name='Calibri', size=9, color='FF808080')
    r = 4
    for k, v in kpi.items():
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='生徒').font = BOLD
    ws.cell(row=r, column=2, value='科目').font = BOLD
    ws.cell(row=r, column=3, value='配置').font = BOLD
    ws.cell(row=r, column=4, value='必要').font = BOLD
    ws.cell(row=r, column=5, value='不足').font = BOLD
    r += 1
    for name, subj, p, c in report:
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=subj)
        ws.cell(row=r, column=3, value=p)
        ws.cell(row=r, column=4, value=c)
        if p < c:
            cc = ws.cell(row=r, column=5, value=c - p)
            cc.fill = PatternFill('solid', fgColor='FFFFC7CE')
        r += 1
    for col, w in zip('ABCDE', (16, 8, 8, 8, 8)):
        ws.column_dimensions[col].width = w

    # ---- 生徒別シート（生徒 → 日付 → 時間 の順に並べ替え） ----
    ws = wb.create_sheet('生徒別', 1)
    heads = ['生徒', '学年', '日付', '曜日', '時間', '科目', '講師', '形態']
    for j, h in enumerate(heads, 1):
        cc = ws.cell(row=1, column=j, value=h)
        cc.font, cc.alignment, cc.fill, cc.border = BOLD, CEN, HEAD, BORDER
    rows = {n: [] for n in bd.students}
    di = {d['sheet']: i for i, d in enumerate(bd.dates)}
    for (sh, t, b), lst in bd.seat.items():
        for e in lst:
            rows[e['name']].append((di[sh], b, sh, t, e, len(lst)))
    r = 2
    for nm, lst in rows.items():
        if not lst:
            continue
        lst.sort(key=lambda x: (x[0], x[1]))
        top = r
        for i, b, sh, t, e, n in lst:
            d = bd.byd[sh]
            vals = [nm, e['grade'], d['ds'], d['wd'], SLOTNAME[b], e['subj'], t,
                    '1対2' if n == SEATS else '1対1']
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.font, c.alignment, c.border = FONT, CEN, BORDER
            r += 1
        ws.merge_cells(start_row=top, start_column=1, end_row=r - 1, end_column=1)
        ws.merge_cells(start_row=top, start_column=2, end_row=r - 1, end_column=2)
    ws.freeze_panes = 'A2'
    for col, w in zip('ABCDEFGH', (14, 6, 8, 6, 8, 8, 12, 8)):
        ws.column_dimensions[col].width = w

    wb.save(OUT)


# ============================================================ main
def main():
    dates, teachers, av, students, order = load_all()
    bd = Board(dates, teachers, av, students)
    for nm, st in students.items():
        st['_rule'] = parse_ng(st['ng'])

    tasks = []
    for nm in order:
        st = students[nm]
        for subj, cnt, tl, mode in st['subs']:
            tasks.append({'name': nm, 'grade': st['grade'], 'cap': st['cap'],
                          'rule': st['_rule'], 'subj': subj, 'count': cnt,
                          'teachers': tl, 'mode': mode, 'placed': 0})
    for t in tasks:
        t['slack'] = feasible_slots(bd, t) - t['count']

    # MRV: 余裕の小さい順 → 専任優先 → 多コマ優先
    for t in sorted(tasks, key=lambda x: (x['slack'], 0 if len(x['teachers']) == 1 else 1, -x['count'])):
        for i in range(t['count']):
            tl = t['teachers'] if t['mode'] == '順' else \
                t['teachers'][i % len(t['teachers']):] + t['teachers'][:i % len(t['teachers'])]
            if any(place_one(bd, t, tt) for tt in tl):
                t['placed'] += 1

    m1 = pair_pass(bd)
    g1 = gap_pass(bd)
    m2 = pair_pass(bd)

    # ---- KPI ----
    total_placed = sum(len(v) for v in bd.seat.values())
    paired = sum(len(v) for v in bd.seat.values() if len(v) == SEATS)
    gaps = sum(1 for nm in students for d in dates
               if not contiguous(bd.busy[nm][d['sheet']]))
    need = sum(t['count'] for t in tasks)
    got = sum(t['placed'] for t in tasks)
    open_seats = sum(SEATS for d in dates for t in teachers for b in range(5)
                     if bd.avail(t, d['ds'], b))
    kpi = {
        '必要コマ数': need,
        '配置できたコマ数': got,
        '充足率': f'{got/need*100:.1f}%',
        '1対2率（相方がいるコマの割合）': f'{paired/total_placed*100:.1f}%' if total_placed else '-',
        '生徒の空きコマが出た「日」の数': gaps,
        '講師の座席稼働率': f'{total_placed/open_seats*100:.1f}%' if open_seats else '-',
        '改善パス（1対2への統合／空きコマ解消）': f'{m1+m2}件 / {g1}件',
    }
    report = [(t['name'], t['subj'], t['placed'], t['count']) for t in tasks]
    write_excel(bd, kpi, report)

    print('=' * 52)
    for k, v in kpi.items():
        print(f'  {k:<34}: {v}')
    print('=' * 52)
    short = [r for r in report if r[2] < r[3]]
    if short:
        print('★ 物理的に入りきらなかった科目（NGを緩めるか担当を増やす必要あり）:')
        for n, s, p, c in short:
            print(f'   {n} / {s}  {p}/{c}')
    else:
        print('★ 不足なし（全コマ配置完了）')
    print(f'\n出力: {OUT}')


if __name__ == '__main__':
    main()
