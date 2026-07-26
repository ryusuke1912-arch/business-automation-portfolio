# -*- coding: utf-8 -*-
"""
デモ用の見積書ファイル群を input/ に生成する。
現場を再現するため、あえて次のバラつきを混ぜている：
  - レイアウト3種（項目のセル位置・ラベル文言が違う）
  - 表記ゆれ（株式会社／(株)、全角数字、「円」付き金額、カンマ区切り）
  - 消費税行が無いファイル、明細行数がバラバラ
※ 社名・担当者名はすべて架空です。
"""
import os, random, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

random.seed(4242)
HERE = os.path.dirname(os.path.abspath(__file__))
IN = os.path.join(HERE, 'input')
os.makedirs(IN, exist_ok=True)
for fn in os.listdir(IN):
    if fn.endswith('.xlsx'):
        os.remove(os.path.join(IN, fn))

CLIENTS = [
    ('株式会社ミナトデザイン', '(株)ミナトデザイン'),
    ('有限会社アオイ工業', '(有)アオイ工業'),
    ('株式会社ハルカゼ商会', '株式会社ハルカゼ商会'),
    ('カワセ電機株式会社', 'カワセ電機(株)'),
    ('株式会社ノガミ物産', '(株)ノガミ物産'),
    ('株式会社トキワソフト', '株式会社トキワソフト'),
]
STAFF = ['大村 健一', '志賀 優子', '成田 亮', '本橋 さやか']
ITEMS = [
    ('Webサイト制作一式', 180000), ('ロゴデザイン', 60000), ('撮影・レタッチ', 45000),
    ('サーバ初期構築', 35000), ('保守サポート（月額）', 15000), ('原稿ライティング', 28000),
    ('印刷物デザイン', 52000), ('システム改修', 120000), ('データ移行', 75000),
    ('操作研修（半日）', 40000),
]
TITLES = ['コーポレートサイト リニューアル', '会社案内パンフレット制作',
          '基幹システム改修', 'ECサイト構築', '採用サイト制作', '業務システム保守']


def zen(s):
    """半角→全角（表記ゆれ再現用）"""
    return str(s).translate(str.maketrans('0123456789-', '０１２３４５６７８９－'))


def money(v, style):
    if style == 0:
        return v                      # 数値そのまま
    if style == 1:
        return f'{v:,}'               # "180,000"
    return f'{v:,}円'                 # "180,000円"


BOLD = Font(bold=True)
RIGHT = Alignment(horizontal='right')

base = datetime.date(2026, 1, 15)
count = 0
for i in range(24):
    variant = i % 3
    formal, short = random.choice(CLIENTS)
    name = random.choice([formal, short])
    date = base + datetime.timedelta(days=random.randint(0, 180))
    no = f'Q-2026-{1001+i}'
    title = random.choice(TITLES)
    staff = random.choice(STAFF)
    mstyle = random.choice([0, 0, 1, 2])
    lines = random.sample(ITEMS, random.randint(2, 5))
    rows = [(nm, random.randint(1, 3), up) for nm, up in lines]
    sub = sum(q * up for _, q, up in rows)
    tax = int(sub * 0.1)
    has_tax = random.random() > 0.2                # 2割は税行なし（税込表記）

    wb = Workbook()
    ws = wb.active
    ws.title = '見積書'

    if variant == 0:
        ws['A1'] = '御 見 積 書'; ws['A1'].font = Font(size=16, bold=True)
        ws['A3'] = '見積番号'; ws['B3'] = no
        ws['A4'] = '発行日';   ws['B4'] = date.strftime('%Y/%m/%d')
        ws['A6'] = '御社名';   ws['B6'] = name + ' 御中'
        ws['A7'] = '件名';     ws['B7'] = title
        ws['A8'] = '担当者';   ws['B8'] = staff
        hr = 10
    elif variant == 1:
        ws['A1'] = '見積書'; ws['A1'].font = Font(size=16, bold=True)
        ws['D2'] = 'No.';    ws['E2'] = no
        ws['D3'] = '日付';   ws['E3'] = zen(date.strftime('%Y-%m-%d'))
        ws['A4'] = name + ' 御中'; ws['A4'].font = BOLD
        ws['A6'] = '件　名'; ws['B6'] = title
        ws['D6'] = '営業担当'; ws['E6'] = staff
        hr = 9
    else:
        ws['A2'] = 'お見積り'; ws['A2'].font = Font(size=14, bold=True)
        ws['A4'] = '宛先';       ws['C4'] = name
        ws['A5'] = '見積No';     ws['C5'] = no
        ws['A6'] = '見積日';     ws['C6'] = date.strftime('%Y年%m月%d日')
        ws['A7'] = '案件名';     ws['C7'] = title
        ws['A8'] = '担当';       ws['C8'] = staff
        hr = 11

    heads = ['品名', '数量', '単価', '金額'] if variant != 2 else ['内容', '数量', '単価', '小計']
    for j, h in enumerate(heads):
        c = ws.cell(row=hr, column=1 + j, value=h)
        c.font = BOLD
    r = hr + 1
    for nm, q, up in rows:
        ws.cell(row=r, column=1, value=nm)
        ws.cell(row=r, column=2, value=zen(q) if variant == 1 else q)
        ws.cell(row=r, column=3, value=money(up, mstyle)).alignment = RIGHT
        ws.cell(row=r, column=4, value=money(q * up, mstyle)).alignment = RIGHT
        r += 1
    r += 1
    ws.cell(row=r, column=3, value='小計').font = BOLD
    ws.cell(row=r, column=4, value=money(sub, mstyle))
    r += 1
    if has_tax:
        ws.cell(row=r, column=3, value='消費税').font = BOLD
        ws.cell(row=r, column=4, value=money(tax, mstyle))
        r += 1
        total = sub + tax
    else:
        total = sub
    ws.cell(row=r, column=3, value='合計金額' if variant else '合計').font = BOLD
    ws.cell(row=r, column=4, value=money(total, mstyle)).font = BOLD

    for col, w in zip('ABCDE', (26, 8, 14, 14, 16)):
        ws.column_dimensions[col].width = w

    fn = f'見積書_{no}_{name.replace("株式会社","").replace("有限会社","").replace("(株)","").replace("(有)","")}.xlsx'
    wb.save(os.path.join(IN, fn))
    count += 1

print(f'{count}件の見積書を input/ に生成しました（レイアウト3種・表記ゆれ入り）')
