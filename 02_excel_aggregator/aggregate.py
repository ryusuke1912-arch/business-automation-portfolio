# -*- coding: utf-8 -*-
"""
バラバラな書式のExcel見積書を一括で読み取り、一覧表・集計表に変換するツール。

入力 : input/*.xlsx （レイアウトも表記もバラバラでよい）
出力 : 見積一覧.xlsx （一覧／明細／取引先別／月別／要確認 の5シート）
実行 : python aggregate.py

■ 設計方針：セル位置を決め打ちしない
  この種の集計が破綻する原因は「B3が見積番号」のような位置固定です。
  取引先ごと・年度ごとにフォーマットが変わると即座に動かなくなります。
  本ツールは "ラベル文字列を探して、その隣の値を取る" 方式で実装しており、
  項目が別のセルへ動いても、ラベルの言い回しが違っても読み取れます。

■ 吸収している「ゆれ」
  - 項目位置のずれ（レイアウト3種）
  - ラベル文言の違い（見積番号 / No. / 見積No、発行日 / 日付 / 見積日 …）
  - 全角数字、カンマ区切り、「円」付き金額、「2026年01月15日」形式の日付
  - 会社名の表記ゆれ（株式会社 / (株) / ㈱）→ 正規化して名寄せ
  - 消費税行が無いファイル（税込表記）
  - 明細の行数がファイルごとに違う
"""
import os, re, glob, time, unicodedata, datetime
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
INDIR = os.path.join(HERE, 'input')
OUT = os.path.join(HERE, '見積一覧.xlsx')

# ---- 探すラベル（先に書いたものを優先） ----
L_NO    = ['見積番号', '見積no', 'no.', 'no', '番号']
L_DATE  = ['発行日', '見積日', '日付', '作成日']
L_CLI   = ['御社名', '宛先', 'お客様名', '得意先', '取引先', '会社名']
L_TITLE = ['件名', '案件名', '件　名', '工事名', '内容']
L_STAFF = ['担当者', '営業担当', '担当']
L_SUB   = ['小計']
L_TAX   = ['消費税', '税額', '内税', '外税']
L_TOTAL = ['合計金額', '合計', '総額', 'お見積金額']


def nz(v):
    """セル値を正規化した文字列に（全角→半角、空白除去）"""
    if v is None:
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y/%m/%d')
    return unicodedata.normalize('NFKC', str(v)).replace(' ', '').replace('　', '').strip()


def to_int(v):
    """'180,000円' / '１８００００' / 180000 → 180000 （不可ならNone）"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = nz(v)
    s = re.sub(r'[^\d\-]', '', s)
    return int(s) if s not in ('', '-') else None


def to_date(v):
    """複数形式の日付を date に。失敗したら None"""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.date() if isinstance(v, datetime.datetime) else v
    s = nz(v)
    m = re.search(r'(\d{4})[/\-年](\d{1,2})[/\-月](\d{1,2})', s)
    if not m:
        return None
    try:
        return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def norm_client(s):
    """会社名の表記ゆれを正規化して名寄せする"""
    s = nz(s).replace('御中', '')
    s = s.replace('㈱', '株式会社').replace('(株)', '株式会社')
    s = s.replace('㈲', '有限会社').replace('(有)', '有限会社')
    return s.strip()


class Sheet:
    """ラベル探索のためにシート全体を一度スキャンして持つ"""

    def __init__(self, ws):
        self.ws = ws
        self.cells = {}          # (row, col) -> 正規化文字列
        self.raw = {}            # (row, col) -> 生の値
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    self.cells[(c.row, c.column)] = nz(c.value).lower()
                    self.raw[(c.row, c.column)] = c.value
        self.maxcol = ws.max_column or 1

    def find_label(self, labels, min_row=None):
        """ラベルに完全一致するセルを探す（先頭の候補ほど優先）
        min_row を指定すると、その行以降だけを探す。
        （「小計」は明細の列見出しとしても現れるため、集計行だけを見たいときに使う）"""
        for lb in labels:
            key = nz(lb).lower()
            for pos in sorted(self.cells):
                if min_row is not None and pos[0] < min_row:
                    continue
                if self.cells[pos] == key:
                    return pos
        return None

    def value_near(self, pos, maxspan=4):
        """ラベルの右→下の順に、最初に見つかった非空セルを値とみなす"""
        r, c = pos
        for cc in range(c + 1, min(c + 1 + maxspan, self.maxcol + 2)):
            if self.raw.get((r, cc)) not in (None, ''):
                return self.raw[(r, cc)]
        for rr in range(r + 1, r + 3):
            if self.raw.get((rr, c)) not in (None, ''):
                return self.raw[(rr, c)]
        return None

    def get(self, labels, min_row=None):
        pos = self.find_label(labels, min_row)
        return self.value_near(pos) if pos else None

    def find_suffix(self, suffix):
        """'○○ 御中' のようにラベルが無く語尾で判別する項目用"""
        key = nz(suffix)
        for pos, txt in self.cells.items():
            if txt.endswith(key) and len(txt) > len(key):
                return self.raw[pos]
        return None

    def detail_block(self):
        """明細のヘッダ行を検出し、以降の行を読む。
        戻り値: (明細リスト, ヘッダ行, 明細の最終行)"""
        head = None
        for (r, c), txt in sorted(self.cells.items()):
            if txt in ('品名', '内容', '品目', '項目'):
                rowtxt = [self.cells.get((r, cc), '') for cc in range(1, self.maxcol + 1)]
                if '数量' in rowtxt and '単価' in rowtxt:
                    head = (r, rowtxt)
                    break
        if not head:
            return [], None, None
        r0, rowtxt = head
        col = {}
        for i, t in enumerate(rowtxt):
            if t in ('品名', '内容', '品目', '項目'): col['name'] = i + 1
            elif t == '数量':                        col['qty'] = i + 1
            elif t == '単価':                        col['price'] = i + 1
            elif t in ('金額', '小計'):              col['amount'] = i + 1
        out = []
        r = r0 + 1
        blank = 0
        last = r0
        while blank < 2 and r < r0 + 60:
            nm = self.raw.get((r, col.get('name', 1)))
            if nm in (None, ''):
                blank += 1
                r += 1
                continue
            blank = 0
            out.append({
                '品名': nz(nm),
                '数量': to_int(self.raw.get((r, col.get('qty', 0)))),
                '単価': to_int(self.raw.get((r, col.get('price', 0)))),
                '金額': to_int(self.raw.get((r, col.get('amount', 0)))),
            })
            last = r
            r += 1
        return out, r0, last


def parse_file(path):
    wb = load_workbook(path, data_only=True)
    sh = Sheet(wb[wb.sheetnames[0]])

    # 明細を先に確定させ、集計行（小計・税・合計）は明細より下だけを探す。
    # そうしないと明細の列見出し「小計」をつかんでしまう。
    details, dhead, dlast = sh.detail_block()
    below = (dlast + 1) if dlast else None

    client = sh.get(L_CLI) or sh.find_suffix('御中')
    rec = {
        'ファイル名': os.path.basename(path),
        '見積番号': nz(sh.get(L_NO)),
        '発行日': to_date(sh.get(L_DATE)),
        '取引先': norm_client(client) if client else '',
        '件名': nz(sh.get(L_TITLE)),
        '担当': nz(sh.get(L_STAFF)),
        '小計': to_int(sh.get(L_SUB, below)),
        '消費税': to_int(sh.get(L_TAX, below)),
        '合計': to_int(sh.get(L_TOTAL, below)),
    }
    rec['明細件数'] = len(details)

    # ---- 検算：明細合計と小計が一致するか ----
    warn = []
    dsum = sum(d['金額'] for d in details if d['金額'] is not None)
    if rec['小計'] is not None and dsum and dsum != rec['小計']:
        warn.append(f'明細合計{dsum:,}≠小計{rec["小計"]:,}')
    if rec['合計'] is not None and rec['小計'] is not None:
        tax = rec['消費税'] or 0
        if rec['小計'] + tax != rec['合計']:
            warn.append('小計+税≠合計')
    if rec['消費税'] is None:
        warn.append('消費税行なし（税込とみなす）')
    for k in ('見積番号', '発行日', '取引先', '合計'):
        if not rec[k]:
            warn.append(f'{k}を抽出できず')
    rec['要確認'] = ' / '.join(warn)
    return rec, details


# ============================================================ 出力
BOLD = Font(bold=True)
HEAD = PatternFill('solid', fgColor='FFD9E1F2')
WARN = PatternFill('solid', fgColor='FFFFF2CC')
thin = Side(style='thin', color='FFBFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
YEN = '#,##0'


def put_table(ws, cols, rows, widths=None):
    for j, c in enumerate(cols, start=1):
        cc = ws.cell(row=1, column=j, value=c)
        cc.font, cc.fill, cc.border = BOLD, HEAD, BD
        cc.alignment = Alignment(horizontal='center')
    for i, r in enumerate(rows, start=2):
        for j, c in enumerate(cols, start=1):
            v = r.get(c)
            cc = ws.cell(row=i, column=j, value=v)
            cc.border = BD
            if isinstance(v, int) and c not in ('明細件数', '数量', '件数'):
                cc.number_format = YEN
            if isinstance(v, datetime.date):
                cc.number_format = 'yyyy/mm/dd'
        if r.get('要確認'):
            for j in range(1, len(cols) + 1):
                ws.cell(row=i, column=j).fill = WARN
    ws.freeze_panes = 'A2'
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[chr(64 + j)].width = w


def main():
    t0 = time.time()
    files = sorted(glob.glob(os.path.join(INDIR, '*.xlsx')))
    if not files:
        print('input/ に .xlsx がありません。先に make_demo_data.py を実行してください。')
        return

    recs, alldet = [], []
    for p in files:
        try:
            rec, det = parse_file(p)
        except Exception as e:                       # 1件壊れても全体を止めない
            recs.append({'ファイル名': os.path.basename(p), '要確認': f'読込エラー: {e}'})
            continue
        recs.append(rec)
        for d in det:
            d.update({'見積番号': rec['見積番号'], '取引先': rec['取引先'],
                      'ファイル名': rec['ファイル名']})
            alldet.append(d)

    wb = Workbook()

    ws = wb.active
    ws.title = '一覧'
    cols = ['ファイル名', '見積番号', '発行日', '取引先', '件名', '担当',
            '明細件数', '小計', '消費税', '合計', '要確認']
    put_table(ws, cols, recs, [34, 13, 12, 22, 26, 11, 9, 12, 11, 13, 34])

    ws = wb.create_sheet('明細')
    put_table(ws, ['見積番号', '取引先', '品名', '数量', '単価', '金額', 'ファイル名'],
              alldet, [13, 22, 26, 8, 12, 12, 34])

    # 取引先別
    agg = {}
    for r in recs:
        c = r.get('取引先') or '(不明)'
        a = agg.setdefault(c, {'取引先': c, '件数': 0, '合計': 0})
        a['件数'] += 1
        a['合計'] += r.get('合計') or 0
    rows = sorted(agg.values(), key=lambda x: -x['合計'])
    for a in rows:
        a['平均単価'] = a['合計'] // a['件数'] if a['件数'] else 0
    ws = wb.create_sheet('取引先別')
    put_table(ws, ['取引先', '件数', '合計', '平均単価'], rows, [24, 8, 14, 14])

    # 月別
    mag = {}
    for r in recs:
        d = r.get('発行日')
        k = d.strftime('%Y/%m') if d else '(不明)'
        a = mag.setdefault(k, {'年月': k, '件数': 0, '合計': 0})
        a['件数'] += 1
        a['合計'] += r.get('合計') or 0
    ws = wb.create_sheet('月別')
    put_table(ws, ['年月', '件数', '合計'], sorted(mag.values(), key=lambda x: x['年月']),
              [12, 8, 14])

    ws = wb.create_sheet('要確認')
    put_table(ws, ['ファイル名', '見積番号', '取引先', '要確認'],
              [r for r in recs if r.get('要確認')], [34, 13, 22, 46])

    wb.save(OUT)

    ok = sum(1 for r in recs if not r.get('要確認'))
    total = sum(r.get('合計') or 0 for r in recs)
    fields = ['見積番号', '発行日', '取引先', '件名', '担当', '合計']
    filled = sum(1 for r in recs for f in fields if r.get(f))
    print('=' * 52)
    print(f'  処理ファイル数        : {len(files)} 件')
    print(f'  項目の抽出成功率      : {filled/(len(recs)*len(fields))*100:.1f}%  '
          f'({filled}/{len(recs)*len(fields)} 項目)')
    print(f'  警告なしで通った件数  : {ok} / {len(recs)} 件')
    print(f'  抽出した明細行        : {len(alldet)} 行')
    print(f'  合計金額              : {total:,} 円')
    print(f'  処理時間              : {time.time()-t0:.2f} 秒')
    print('=' * 52)
    for r in recs:
        if r.get('要確認'):
            print(f'  [要確認] {r["ファイル名"]}  … {r["要確認"]}')
    print(f'\n出力: {OUT}')


if __name__ == '__main__':
    main()
